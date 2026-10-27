#!/usr/bin/env python3
"""
fill_actuals.py -- fill Sewing Actuals into a previously-generated
Plan vs Actuals workbook, using a "linewise performance report".

Usage:
    python3 fill_actuals.py <generated_workbook.xlsx> <performance_report.xlsx> <output.xlsx>
"""
import sys
import json
import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment


def fail(message):
    print(json.dumps({"error": message}), file=sys.stderr)
    sys.exit(1)


def parse_report_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str):
        v_str = v.strip()
        for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d-%b-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%y", "%Y/%m/%d"):
            try:
                return datetime.datetime.strptime(v_str, fmt).date()
            except ValueError:
                continue
    return None


def load_excel_workbook(filepath, data_only=True):
    try:
        return openpyxl.load_workbook(filepath, data_only=data_only)
    except Exception as e:
        err_msg = str(e).lower()
        if "xls" in err_msg or filepath.lower().endswith(".xls"):
            try:
                import xlrd
            except ImportError:
                fail("Could not read .xls file: xlrd module is not installed.")

            try:
                wb_xls = xlrd.open_workbook(filepath)
            except Exception as ex_xls:
                fail(f"Could not open .xls file: {ex_xls}")

            wb_out = openpyxl.Workbook()
            wb_out.remove(wb_out.active)

            for sheet_name in wb_xls.sheet_names():
                sheet_xls = wb_xls.sheet_by_name(sheet_name)
                ws_out = wb_out.create_sheet(title=sheet_name)
                for r in range(sheet_xls.nrows):
                    row_vals = []
                    for c in range(sheet_xls.ncols):
                        cell_val = sheet_xls.cell_value(r, c)
                        cell_type = sheet_xls.cell_type(r, c)
                        if cell_type == xlrd.XL_CELL_DATE:
                            try:
                                dt_tuple = xlrd.xldate_as_tuple(cell_val, wb_xls.datemode)
                                if dt_tuple[0] == 0:
                                    cell_val = datetime.time(*dt_tuple[3:])
                                else:
                                    cell_val = datetime.datetime(*dt_tuple)
                            except Exception:
                                pass
                        elif cell_type in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            cell_val = None
                        elif cell_type == xlrd.XL_CELL_NUMBER:
                            if isinstance(cell_val, float) and cell_val.is_integer():
                                cell_val = int(cell_val)
                        row_vals.append(cell_val)
                    ws_out.append(row_vals)
            return wb_out
        else:
            raise e


def main():
    if len(sys.argv) < 3:
        fail("Usage: fill_actuals.py <generated_workbook.xlsx> <performance_report.xlsx> <output.xlsx>")

    workbook_path = sys.argv[1]
    report_path = sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) >= 4 else "output.xlsx"

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except Exception as e:
        fail(f"Could not open the generated workbook: {e}")

    try:
        wb_perf = load_excel_workbook(report_path, data_only=True)
    except Exception as e:
        fail(f"Could not open the performance report: {e}")

    ws_perf = wb_perf.worksheets[0]
    perf_rows = list(ws_perf.iter_rows(values_only=True))
    if not perf_rows:
        fail("Performance report sheet is empty.")

    perf_header = [str(h).strip().upper() if h else h for h in perf_rows[0]]

    def find_col(*names):
        for name in names:
            if name in perf_header:
                return perf_header.index(name)
        return None

    date_idx = find_col("DATE")
    style_idx = find_col("STYLE NO", "STYLE")
    prod_idx = find_col("PRODUCTION")

    if date_idx is None or style_idx is None or prod_idx is None:
        fail("Performance report is missing DATE, STYLE NO, or PRODUCTION columns.")

    sheet_names = wb.sheetnames

    combined_bases = set()
    for name in sheet_names:
        if name.endswith("-Jacket"):
            base = name[:-len("-Jacket")]
            if f"{base}-Pant" in sheet_names:
                combined_bases.add(base)

    def resolve_targets(raw_style):
        if raw_style in sheet_names:
            return [("direct", raw_style)]
        if raw_style in combined_bases:
            return [("combined", raw_style)]
        for sn in sheet_names:
            if raw_style == sn or raw_style.startswith(sn + "-"):
                return [("direct", sn)]
        return []

    sheet_daily = defaultdict(lambda: defaultdict(int))
    combined_daily = defaultdict(lambda: defaultdict(int))
    unmatched_styles = set()

    # Every date that actually appears in the uploaded performance report,
    # regardless of which style it belongs to. This is what tells us the
    # upload "covers" that date -- so any style with no row on that date
    # should get a real 0, not a blank. Dates that never appear here were
    # simply not uploaded yet, and must stay blank.
    uploaded_dates = set()

    for row in perf_rows[1:]:
        raw_style = row[style_idx]
        if not raw_style:
            continue
        raw_style = str(raw_style).strip()
        d = parse_report_date(row[date_idx])
        if d is None:
            continue
        uploaded_dates.add(d)
        prod = row[prod_idx] or 0

        targets = resolve_targets(raw_style)
        if not targets:
            unmatched_styles.add(raw_style)
            continue

        for kind, key in targets:
            if kind == "direct":
                sheet_daily[key][d] += prod
            else:
                combined_daily[key][d] += prod

    for base, by_date in combined_daily.items():
        jacket_sheet = f"{base}-Jacket"
        pant_sheet = f"{base}-Pant"
        for d, total in by_date.items():
            jacket_qty = round(total / 2)
            pant_qty = total - jacket_qty
            sheet_daily[jacket_sheet][d] += jacket_qty
            sheet_daily[pant_sheet][d] += pant_qty

    sheets_updated = 0
    cells_filled = 0
    center_align = Alignment(horizontal='center', vertical='center')

    valid_uploaded_dates = [ud for ud in uploaded_dates if isinstance(ud, datetime.date)]
    max_uploaded_date = max(valid_uploaded_dates) if valid_uploaded_dates else None

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        daily = sheet_daily.get(sheet_name)
        touched = False
        last_date_row = 4

        for r in range(4, ws.max_row + 1):
            date_val = ws.cell(row=r, column=3).value
            if date_val is None:
                continue

            last_date_row = r
            parsed_d = parse_report_date(date_val)
            d = parsed_d if parsed_d is not None else (date_val.date() if hasattr(date_val, "date") else date_val)

            # Sewing Actuals = column F (Col 6)
            f_cell = ws.cell(row=r, column=6)
            f_cell.alignment = center_align

            if d in uploaded_dates:
                # This date was covered by the uploaded report. Use the
                # matched production qty if we found one for this style;
                # otherwise this style genuinely had zero production that
                # day, so record a real 0 instead of leaving it blank.
                f_cell.value = daily[d] if (daily and d in daily) else 0
                cells_filled += 1
                touched = True
            elif max_uploaded_date and isinstance(d, datetime.date) and d > max_uploaded_date:
                # Date is past the maximum uploaded date -> clear actuals
                f_cell.value = None

            # Preserve & re-assert exact IF(ISNUMBER(...)) formulas across ALL 3 blocks for every row 4..max_row
            ws.cell(row=r, column=5, value=f'=IF(ISNUMBER(F{r}),SUM($D$4:D{r}),"")')     # E: Sew Cum Plan
            ws.cell(row=r, column=7, value=f'=IF(ISNUMBER(F{r}),SUM($F$4:F{r}),"")')     # G: Sew Cum Actual
            ws.cell(row=r, column=8, value=f'=IF(ISNUMBER(G{r}),N(E{r})-G{r},"")')      # H: Sew Deviation (blank Cum Plan treated as 0)
            ws.cell(row=r, column=10, value=f'=IF(ISNUMBER(K{r}),SUM($I$4:I{r}),"")')   # J: Wash Cum Plan
            ws.cell(row=r, column=12, value=f'=IF(ISNUMBER(K{r}),SUM($K$4:K{r}),"")')   # L: Wash Cum Actual
            ws.cell(row=r, column=13, value=f'=IF(ISNUMBER(L{r}),N(J{r})-L{r},"")')     # M: Wash Deviation (blank Cum Plan treated as 0)
            ws.cell(row=r, column=15, value=f'=IF(ISNUMBER(P{r}),SUM($N$4:N{r}),"")')   # O: Pack Cum Plan
            ws.cell(row=r, column=17, value=f'=IF(ISNUMBER(P{r}),SUM($P$4:P{r}),"")')   # Q: Pack Cum Actual
            ws.cell(row=r, column=18, value=f'=IF(ISNUMBER(Q{r}),N(O{r})-Q{r},"")')     # R: Pack Deviation (blank Cum Plan treated as 0)

            for col_idx in (5, 7, 8, 10, 12, 13, 15, 17, 18):
                ws.cell(row=r, column=col_idx).alignment = center_align

        # NOTE: Order Summary columns (T to AB) are intentionally NOT touched here.
        # That section (Style, X-fty, Qty, Sewing completed, Bal to sew, Wash
        # completed, Bal to Wash, Pack completed, Bal to Pack) is maintained
        # manually by the user in the uploaded Existing Workbook and must be
        # returned exactly as uploaded -- this script only fills the Sewing
        # Details / Wash Details / Pack Details block (columns A-R).

        if touched:
            sheets_updated += 1

    try:
        wb.save(out_path)
    except Exception as e:
        fail(f"Could not save output file: {e}")

    print(json.dumps({
        "success": True,
        "sheetsUpdated": sheets_updated,
        "cellsFilled": cells_filled,
        "combinedSplitStyles": sorted(combined_daily.keys()),
        "unmatchedStyles": sorted(unmatched_styles),
    }))


if __name__ == "__main__":
    main()