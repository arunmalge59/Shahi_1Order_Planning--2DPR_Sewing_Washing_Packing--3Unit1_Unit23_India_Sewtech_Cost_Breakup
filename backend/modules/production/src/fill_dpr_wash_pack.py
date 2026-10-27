#!/usr/bin/env python3
"""
fill_dpr_wash_pack.py -- fill Wash & Pack Actuals for Unit 23 into a previously-generated
Plan vs Actuals workbook using a DPR Daily Performance Report (.xls / .xlsx).

Usage:
    python3 fill_dpr_wash_pack.py <generated_workbook.xlsx> <dpr_report.xls/xlsx> <output.xlsx> [report_date_YYYY-MM-DD]
"""
import sys
import os
import re
import json
import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment


def fail(message):
    print(json.dumps({"error": message}), file=sys.stderr)
    sys.exit(1)


MONTH_MAP = {
    'jan': 1, 'january': 1,
    'feb': 2, 'february': 2,
    'mar': 3, 'march': 3,
    'apr': 4, 'april': 4,
    'may': 5,
    'jun': 6, 'june': 6,
    'jul': 7, 'july': 7,
    'aug': 8, 'august': 8,
    'sep': 9, 'sept': 9, 'september': 9,
    'oct': 10, 'october': 10,
    'nov': 11, 'november': 11,
    'dec': 12, 'december': 12
}


def parse_date_from_string(text, ref_year=None):
    if not text:
        return None

    cleaned = os.path.splitext(os.path.basename(str(text)))[0]
    cleaned = re.sub(r'\(\d+\)', '', cleaned).strip()
    y_ref = ref_year or datetime.date.today().year

    # 1. YYYY-MM-DD or YYYY_MM_DD or YYYY.MM.DD
    m = re.search(r'(\d{4})[-_./](\d{1,2})[-_./](\d{1,2})', cleaned)
    if m:
        y, mth, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime.date(y, mth, d)
        except ValueError:
            pass

    # 2. DD-MM-YYYY or DD_MM_YYYY or DD.MM.YYYY
    m = re.search(r'(\d{1,2})[-_./](\d{1,2})[-_./](\d{4})', cleaned)
    if m:
        d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime.date(y, mth, d)
        except ValueError:
            pass

    # 3. DD-Mon-YYYY or DD Mon YYYY or DD_Mon_YYYY
    m = re.search(r'(\d{1,2})[-_.\s]+([A-Za-z]{3,9})[-_.\s]+(\d{2,4})', cleaned)
    if m:
        d_str, m_str, y_str = m.group(1), m.group(2), m.group(3)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            y = int(y_str) if len(y_str) == 4 else 2000 + int(y_str)
            try:
                return datetime.date(y, mth, int(d_str))
            except ValueError:
                pass

    # 4. Mon-DD-YYYY or Mon DD YYYY or Mon_DD_YYYY
    m = re.search(r'([A-Za-z]{3,9})[-_.\s]+(\d{1,2})[-_.\s]+(\d{2,4})', cleaned)
    if m:
        m_str, d_str, y_str = m.group(1), m.group(2), m.group(3)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            y = int(y_str) if len(y_str) == 4 else 2000 + int(y_str)
            try:
                return datetime.date(y, mth, int(d_str))
            except ValueError:
                pass

    # 5. DD-Mon or DD Mon or DD_Mon
    m = re.search(r'(\d{1,2})[-_.\s]*([A-Za-z]{3,9})(?=[-_.\s]|$)', cleaned)
    if m:
        d_str, m_str = m.group(1), m.group(2)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            try:
                return datetime.date(y_ref, mth, int(d_str))
            except ValueError:
                pass

    # 6. Mon-DD or Mon DD or Mon_DD (e.g. SEP 15, SEP-15, DPR SEP 15)
    m = re.search(r'([A-Za-z]{3,9})[-_.\s]*(\d{1,2})(?=[-_.\s]|$)', cleaned)
    if m:
        m_str, d_str = m.group(1), m.group(2)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            try:
                return datetime.date(y_ref, mth, int(d_str))
            except ValueError:
                pass

    # 7. Short DD-MM or DD_MM or DD.MM
    m = re.search(r'(\d{1,2})[-_./](\d{1,2})', cleaned)
    if m:
        d, mth = int(m.group(1)), int(m.group(2))
        try:
            return datetime.date(y_ref, mth, d)
        except ValueError:
            pass

    return None


def parse_number(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        s = val.replace(',', '').replace(' ', '').replace('$', '').strip()
        try:
            return float(s)
        except ValueError:
            return 0.0
    return 0.0


def parse_style_and_item(style_raw):
    """
    Parses strings like "F2625W172A (BZ28)" or "S2725W901B  (CY37)"
    Returns (style_code, item_code)
    """
    if not style_raw:
        return "", ""
    s = str(style_raw).strip()
    m = re.search(r'^(.*?)\s*\(([^)]+)\)\s*$', s)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return s, ""


def load_excel_workbook(filepath, data_only=True):
    try:
        return openpyxl.load_workbook(filepath, data_only=data_only)
    except Exception as e:
        err_msg = str(e).lower()
        if "xls" in err_msg or filepath.lower().endswith(".xls"):
            try:
                import xlrd
            except ImportError:
                fail("Could not read .xls file: xlrd module is not installed.")

            try:
                wb_xls = xlrd.open_workbook(filepath)
            except Exception as ex_xls:
                fail(f"Could not open .xls file: {ex_xls}")

            wb_out = openpyxl.Workbook()
            wb_out.remove(wb_out.active)

            for sheet_name in wb_xls.sheet_names():
                sheet_xls = wb_xls.sheet_by_name(sheet_name)
                ws_out = wb_out.create_sheet(title=sheet_name)
                for r in range(sheet_xls.nrows):
                    row_vals = []
                    for c in range(sheet_xls.ncols):
                        cell_val = sheet_xls.cell_value(r, c)
                        cell_type = sheet_xls.cell_type(r, c)
                        if cell_type == xlrd.XL_CELL_DATE:
                            try:
                                dt_tuple = xlrd.xldate_as_tuple(cell_val, wb_xls.datemode)
                                if dt_tuple[0] == 0:
                                    cell_val = datetime.time(*dt_tuple[3:])
                                else:
                                    cell_val = datetime.datetime(*dt_tuple)
                            except Exception:
                                pass
                        elif cell_type in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            cell_val = None
                        elif cell_type == xlrd.XL_CELL_NUMBER:
                            if isinstance(cell_val, float) and cell_val.is_integer():
                                cell_val = int(cell_val)
                        row_vals.append(cell_val)
                    ws_out.append(row_vals)
            return wb_out
        else:
            raise e


def find_dpr_columns(rows, max_scan=10):
    buyer_col = None
    style_col = None
    po_col = None
    wash_on_date_col = None
    pack_on_date_col = None
    data_start_row = 3

    if not rows:
        fail("Error: DPR report has no rows.")

    max_scan = min(max_scan, len(rows))
    header_matrix = []
    for r_idx in range(max_scan):
        row_vals = [str(cell.value).strip() if cell and cell.value is not None else "" for cell in rows[r_idx]]
        header_matrix.append(row_vals)

    num_cols = max(len(r) for r in header_matrix) if header_matrix else 0

    # 1. Identify Buyer, Style, and PO columns
    for r in range(max_scan):
        for c in range(num_cols):
            if c < len(header_matrix[r]):
                val = header_matrix[r][c].upper().replace(" ", "").replace("_", "")
                if val in ("BUYER", "CUSTOMER") and buyer_col is None:
                    buyer_col = c
                elif ("STYLE" in val or "STYLE#" in val or "STYLENO" in val) and "BUYER" not in val and style_col is None:
                    style_col = c
                elif val in ("PO", "PONO", "BUYERPO", "ORDERNO") and po_col is None:
                    po_col = c

    if style_col is None:
        style_col = 1
    if buyer_col is None:
        buyer_col = 0
    if po_col is None:
        po_col = 2

    # 2. Identify Section Headers explicitly
    wash_sec_start = None
    pack_sec_start = None

    for r in range(max_scan):
        for c in range(num_cols):
            if c < len(header_matrix[r]):
                raw = header_matrix[r][c].upper()
                clean = raw.replace(" ", "").replace("_", "").replace("-", "")

                # WASHING-RECD section check (must NOT match WASHING-SENT)
                is_wash_recd = (
                    (("WASH" in clean or "W/REC" in clean) and any(k in clean for k in ("RECD", "REC", "RECEIVE", "RECEIVED")) and "SENT" not in clean)
                    or ("WASHING-RECD" in raw or "WASHING RECD" in raw or "WASH RECD" in raw or "W/REC" in raw)
                ) and "SENT" not in clean

                if is_wash_recd and wash_sec_start is None:
                    wash_sec_start = c

                # PACKING section check (must NOT match UNPACKED)
                is_packing = ("PACK" in clean or "PACKING" in clean or "PACKED" in clean) and "UNPACK" not in clean

                if is_packing and pack_sec_start is None:
                    pack_sec_start = c

    # 3. Determine top header column boundaries to establish section column ranges
    top_header_cols = set()
    for r in range(min(3, max_scan)):
        for c in range(num_cols):
            if c < len(header_matrix[r]) and header_matrix[r][c]:
                top_header_cols.add(c)
    sorted_top_cols = sorted(list(top_header_cols))

    def get_section_col_range(start_c):
        next_cols = [c for c in sorted_top_cols if c > start_c]
        end_c = (next_cols[0] - 1) if next_cols else (start_c + 3)
        return start_c, max(start_c, end_c)

    # 4. Search for ON DATE column inside WASHING-RECD section range
    if wash_sec_start is not None:
        w_min, w_max = get_section_col_range(wash_sec_start)
        for c in range(w_min, min(w_max + 1, num_cols)):
            for r in range(max_scan):
                if c < len(header_matrix[r]):
                    cv = header_matrix[r][c].upper().replace(" ", "").replace("_", "").replace("-", "")
                    if cv in ("ONDATE", "TODAY", "ON", "DAILY"):
                        wash_on_date_col = c
                        break
            if wash_on_date_col is not None:
                break

    # 5. Search for ON DATE column inside PACKING section range
    if pack_sec_start is not None:
        p_min, p_max = get_section_col_range(pack_sec_start)
        for c in range(p_min, min(p_max + 1, num_cols)):
            for r in range(max_scan):
                if c < len(header_matrix[r]):
                    cv = header_matrix[r][c].upper().replace(" ", "").replace("_", "").replace("-", "")
                    if cv in ("ONDATE", "TODAY", "ON", "DAILY"):
                        pack_on_date_col = c
                        break
            if pack_on_date_col is not None:
                break

    # Strict check: Fail with clear error if section columns are not confidently identified
    if wash_on_date_col is None:
        fail("Error: Could not confidently identify the Washing ON DATE column in the DPR header (WASHING-RECD section missing or child ON DATE column not found).")

    if pack_on_date_col is None:
        fail("Error: Could not confidently identify the Packing ON DATE column in the DPR header (PACKING section missing or child ON DATE column not found).")

    for r in range(max_scan):
        row_str = "".join(header_matrix[r]).upper()
        if "BUYER" in row_str or "STYLE" in row_str or "ON DATE" in row_str:
            data_start_row = max(data_start_row, r + 2)

    return buyer_col, style_col, po_col, wash_on_date_col, pack_on_date_col, data_start_row


def main():
    if len(sys.argv) < 3:
        fail("Usage: fill_dpr_wash_pack.py <generated_workbook.xlsx> <dpr_report.xls/xlsx> <output.xlsx> [report_date]")

    workbook_path = sys.argv[1]
    report_path = sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) >= 4 else "output.xlsx"
    user_report_date_str = sys.argv[4] if len(sys.argv) >= 5 else None

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except Exception as e:
        fail(f"Could not open target workbook: {e}")

    ref_year = None
    for ws in wb.worksheets:
        for r in range(4, min(15, ws.max_row + 1)):
            v = ws.cell(row=r, column=3).value
            if v and hasattr(v, "year"):
                ref_year = v.year
                break
        if ref_year:
            break

    # The explicitly selected date belongs to this report and takes precedence.
    target_date = None
    if user_report_date_str:
        try:
            target_date = datetime.datetime.strptime(user_report_date_str.strip(), "%Y-%m-%d").date()
        except ValueError:
            target_date = parse_date_from_string(user_report_date_str, ref_year)

    if not target_date:
        target_date = parse_date_from_string(report_path, ref_year)

    if not target_date:
        try:
            wb_dpr_check = load_excel_workbook(report_path, data_only=True)
            for ws_check in wb_dpr_check.worksheets:
                target_date = parse_date_from_string(ws_check.title, ref_year)
                if target_date:
                    break
        except Exception:
            pass

    if not target_date:
        filename_only = os.path.basename(report_path)
        fail(f"Could not detect a valid report date from the DPR filename ('{filename_only}'). Please ensure the DPR file is named with a valid report date (e.g. 'DPR SEP 15.xls' or 'DPR 15-09-2026.xls').")

    try:
        wb_dpr = load_excel_workbook(report_path, data_only=True)
    except Exception as e:
        fail(f"Could not open DPR report file: {e}")

    ws_dpr = wb_dpr.worksheets[0]
    rows = list(ws_dpr.iter_rows(values_only=False))

    buyer_col, style_col, po_col, wash_col, pack_col, data_start = find_dpr_columns(rows)

    print(f"[DPR Parser] Report Date: {target_date}, Style Col: {style_col}, PO Col: {po_col}, Wash ON DATE Col: {wash_col}, Pack ON DATE Col: {pack_col}", file=sys.stderr)

    wash_actuals = defaultdict(float)  # style_code -> sum
    pack_actuals = defaultdict(float)  # style_code -> sum
    matched_styles = set()
    unmatched_styles = set()

    sheet_names = wb.sheetnames

    def resolve_targets(raw_style_code):
        if raw_style_code in sheet_names:
            return [("direct", raw_style_code)]
        for sn in sheet_names:
            if "&" in sn:
                parts = [p.strip() for p in sn.split("&")]
                if raw_style_code in parts:
                    return [("direct", sn)]
        for sn in sheet_names:
            if raw_style_code == sn or raw_style_code.startswith(sn + "-") or sn.startswith(raw_style_code + "-") or raw_style_code.startswith(sn + " "):
                return [("direct", sn)]
        return []

    dpr_rows_processed = 0
    current_buyer = ""

    for r_idx in range(data_start - 1, len(rows)):
        r = [cell.value for cell in rows[r_idx]]
        if not r:
            continue

        buyer_val = str(r[buyer_col]).strip() if len(r) > buyer_col and r[buyer_col] is not None else ""
        if buyer_val:
            current_buyer = buyer_val

        style_raw = str(r[style_col]).strip() if len(r) > style_col and r[style_col] is not None else ""
        po_val = str(r[po_col]).strip() if len(r) > po_col and r[po_col] is not None else ""

        # Skip subtotal/grand-total rows (blank style or blank PO or TOTAL in text)
        if not style_raw:
            continue

        if any(term in style_raw.upper() for term in ("TOTAL", "SUBTOTAL", "GRAND TOTAL", "SUM")):
            continue

        if po_col is not None and (not po_val or any(term in po_val.upper() for term in ("TOTAL", "SUBTOTAL", "GRAND TOTAL", "SUM"))):
            continue

        style_code, item_code = parse_style_and_item(style_raw)
        if not style_code:
            continue

        w_val = parse_number(r[wash_col]) if len(r) > wash_col else 0.0
        p_val = parse_number(r[pack_col]) if len(r) > pack_col else 0.0

        targets = resolve_targets(style_code)
        if not targets:
            unmatched_styles.add(style_code)
            continue

        dpr_rows_processed += 1
        for kind, key in targets:
            matched_styles.add(key)
            wash_actuals[key] += w_val
            pack_actuals[key] += p_val

    sheets_updated = 0
    cells_filled = 0
    center_align = Alignment(horizontal='center', vertical='center')

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        touched = False

        for r in range(4, ws.max_row + 1):
            date_val = ws.cell(row=r, column=3).value
            if date_val is None:
                continue

            parsed_d = parse_date_from_string(date_val, ref_year)
            d = parsed_d if parsed_d is not None else (date_val.date() if hasattr(date_val, "date") else date_val)

            if d == target_date:
                w_tot = wash_actuals.get(sheet_name, 0.0)
                p_tot = pack_actuals.get(sheet_name, 0.0)

                k_cell = ws.cell(row=r, column=11, value=int(w_tot) if w_tot.is_integer() else w_tot)
                p_cell = ws.cell(row=r, column=16, value=int(p_tot) if p_tot.is_integer() else p_tot)
                k_cell.alignment = center_align
                p_cell.alignment = center_align
                cells_filled += 2
                touched = True
            elif target_date and isinstance(d, datetime.date) and isinstance(target_date, datetime.date) and d > target_date:
                # Date is past target_date -> clear actuals
                ws.cell(row=r, column=11, value=None)
                ws.cell(row=r, column=16, value=None)

            # Assert standard formulas for every row
            ws.cell(row=r, column=5, value=f'=IF(ISNUMBER(F{r}),SUM($D$4:D{r}),"")')     # E: Sew Cum Plan
            ws.cell(row=r, column=7, value=f'=IF(ISNUMBER(F{r}),SUM($F$4:F{r}),"")')     # G: Sew Cum Actual
            ws.cell(row=r, column=8, value=f'=IF(ISNUMBER(G{r}),N(E{r})-G{r},"")')      # H: Sew Deviation
            ws.cell(row=r, column=10, value=f'=IF(ISNUMBER(K{r}),SUM($I$4:I{r}),"")')   # J: Wash Cum Plan
            ws.cell(row=r, column=12, value=f'=IF(ISNUMBER(K{r}),SUM($K$4:K{r}),"")')   # L: Wash Cum Actual
            ws.cell(row=r, column=13, value=f'=IF(ISNUMBER(L{r}),N(J{r})-L{r},"")')     # M: Wash Deviation
            ws.cell(row=r, column=15, value=f'=IF(ISNUMBER(P{r}),SUM($N$4:N{r}),"")')   # O: Pack Cum Plan
            ws.cell(row=r, column=17, value=f'=IF(ISNUMBER(P{r}),SUM($P$4:P{r}),"")')   # Q: Pack Cum Actual
            ws.cell(row=r, column=18, value=f'=IF(ISNUMBER(Q{r}),N(O{r})-Q{r},"")')     # R: Pack Deviation

            for col_idx in (5, 7, 8, 10, 12, 13, 15, 17, 18):
                ws.cell(row=r, column=col_idx).alignment = center_align

        if touched:
            sheets_updated += 1

    try:
        wb.save(out_path)
    except Exception as e:
        fail(f"Could not save output file: {e}")

    style_totals_output = {}
    for st in matched_styles:
        style_totals_output[st] = {
            "wash": wash_actuals[st],
            "pack": pack_actuals[st]
        }

    print(json.dumps({
        "success": True,
        "targetDate": str(target_date),
        "detectedWashOnDateCol": wash_col,
        "detectedPackOnDateCol": pack_col,
        "detectedStyleCol": style_col,
        "detectedPoCol": po_col,
        "dprRowsProcessed": dpr_rows_processed,
        "sheetsUpdated": sheets_updated,
        "cellsFilled": cells_filled,
        "styleTotals": style_totals_output,
        "unmatchedStyles": sorted(list(unmatched_styles)),
    }))


if __name__ == "__main__":
    main()
