#!/usr/bin/env python3
"""
fill_order_summary.py -- read an Order Sheet Excel file, group orders by (Style, X-fty date)
filtered to the generated workbook's active month/year period, and populate the Order
Summary section (T4:AB...) with live balance formulas (=Completed - Qty).

Usage:
    python3 fill_order_summary.py <generated_workbook.xlsx> <order_sheet.xlsx> <output.xlsx>
"""
import sys
import json
import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, Border, Side


def fail(message):
    print(json.dumps({"error": message}), file=sys.stderr)
    sys.exit(1)


def parse_number(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        s = val.replace(',', '').replace(' ', '').replace('$', '').strip()
        try:
            return float(s)
        except ValueError:
            return 0.0
    return 0.0


def parse_date_value(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)):
        try:
            if 30000 < v < 60000:
                return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=v)).date()
        except Exception:
            pass
        return None
    if isinstance(v, str):
        s = v.strip()
        for fmt in (
            "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y",
            "%d-%b-%Y", "%d-%b-%y", "%d-%m-%y", "%Y/%m/%d",
            "%d.%m.%Y", "%Y.%m.%d"
        ):
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def main():
    if len(sys.argv) != 4:
        fail("Usage: fill_order_summary.py <generated_workbook.xlsx> <order_sheet.xlsx> <output.xlsx>")

    workbook_path, order_sheet_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except Exception as e:
        fail(f"Could not open the generated workbook: {e}")

    try:
        try:
            wb_orders = openpyxl.load_workbook(order_sheet_path, read_only=True, data_only=True)
        except Exception:
            wb_orders = openpyxl.load_workbook(order_sheet_path, data_only=True)
    except Exception as e:
        fail(f"Could not open the order sheet: {e}")

    # Extract active period (year, month) from generated workbook Column C dates
    active_months = set()
    for ws in wb.worksheets:
        for r in range(4, min(50, ws.max_row + 1)):
            val = ws.cell(row=r, column=3).value
            if val is not None:
                d = parse_date_value(val)
                if d is not None:
                    active_months.add((d.year, d.month))

    ws_orders = wb_orders.worksheets[0]
    rows = list(ws_orders.iter_rows(values_only=True))
    if hasattr(wb_orders, "close"):
        try:
            wb_orders.close()
        except Exception:
            pass

    if not rows:
        fail("Order sheet is empty.")

    header_row_idx = None
    for idx, r in enumerate(rows[:20]):
        if not r:
            continue
        cleaned = [str(c).strip().lower() if c is not None else '' for c in r]
        has_style = any(c in ('style', 'style number', 'style no', 'style no.') or ('style' in c and 'buyer' not in c and 'customer' not in c) for c in cleaned)
        has_qty = any(c in ('qty', 'quantity', 'order qty', 'total qty') or 'qty' in c or 'quantity' in c for c in cleaned)
        if has_style and has_qty:
            header_row_idx = idx
            break

    if header_row_idx is None:
        fail("Invalid Order Sheet. Required columns missing. Please upload the correct Order Sheet.")

    header = rows[header_row_idx]
    cleaned_header = [str(c).strip().lower() if c is not None else '' for c in header]

    def find_col(exact_candidates, substring_candidates=None, exclude_substrings=None):
        for candidate in exact_candidates:
            for idx, cell in enumerate(cleaned_header):
                if cell == candidate:
                    return idx
        if substring_candidates:
            for candidate in substring_candidates:
                for idx, cell in enumerate(cleaned_header):
                    if candidate in cell:
                        if exclude_substrings and any(ex in cell for ex in exclude_substrings):
                            continue
                        return idx
        return None

    style_idx = find_col(
        ['style number', 'style no.', 'style no', 'style', 'style_number', 'styleno'],
        ['style'],
        exclude_substrings=['buyer', 'customer']
    )

    qty_idx = find_col(
        ['order qty', 'total qty', 'quantity', 'qty', 'order quantity', 'po qty', 'order_qty'],
        ['qty', 'quantity']
    )

    xfty_idx = find_col(
        ['xfty date', 'x-fty date', 'x-ftydate', 'x-fty', 'xfty', 'x-factory', 'xfactory', 'ex-factory', 'crd', 'pcd'],
        ['xfty', 'x-fty', 'x-factory', 'xfactory', 'ex-factory', 'crd', 'pcd', 'ship date', 'del date', 'delivery']
    )

    sewing_idx = find_col(
        ['sewing completed', 'sewing complete', 'sewing qty', 'sewing'],
        ['sewing']
    )

    wash_idx = find_col(
        ['wash completed', 'wash complete', 'wash qty', 'wash'],
        ['wash']
    )

    pack_idx = find_col(
        ['pack completed', 'pack complete', 'pack qty', 'pack', 'packing'],
        ['pack']
    )

    if style_idx is None or qty_idx is None or xfty_idx is None:
        fail("Invalid Order Sheet. Required columns missing. Please upload the correct Order Sheet.")

    # Group by (Style, X-fty date), filtering X-fty date to active_months
    style_xfty_groups = defaultdict(lambda: defaultdict(lambda: {'qty': 0.0, 'sewing': 0.0, 'wash': 0.0, 'pack': 0.0}))
    order_styles_set = set()

    for r in rows[header_row_idx + 1:]:
        if not r or len(r) <= style_idx or r[style_idx] is None:
            continue
        raw_style = str(r[style_idx]).strip()
        if not raw_style or raw_style.lower() in ('total', 'grand total', 'sum', 'style'):
            continue

        xfty_val = parse_date_value(r[xfty_idx]) if (len(r) > xfty_idx) else None

        # Period filter: if active_months is determined, only include X-fty dates falling within active_months
        if xfty_val is not None and active_months:
            if (xfty_val.year, xfty_val.month) not in active_months:
                continue

        qty_val = parse_number(r[qty_idx]) if (len(r) > qty_idx) else 0.0
        sewing_val = parse_number(r[sewing_idx]) if (sewing_idx is not None and len(r) > sewing_idx) else 0.0
        wash_val = parse_number(r[wash_idx]) if (wash_idx is not None and len(r) > wash_idx) else 0.0
        pack_val = parse_number(r[pack_idx]) if (pack_idx is not None and len(r) > pack_idx) else 0.0

        entry = style_xfty_groups[raw_style][xfty_val]
        entry['qty'] += qty_val
        entry['sewing'] += sewing_val
        entry['wash'] += wash_val
        entry['pack'] += pack_val
        order_styles_set.add(raw_style)

    sheet_names = wb.sheetnames

    def find_order_key(sheet_name):
        if sheet_name in style_xfty_groups:
            return sheet_name
        for order_style in style_xfty_groups:
            if order_style == sheet_name or order_style.startswith(sheet_name + "-") or sheet_name.startswith(order_style + "-") or order_style.startswith(sheet_name + " ") or sheet_name.startswith(order_style + " "):
                return order_style
        sheet_base = sheet_name.split('-')[0].split(' ')[0]
        for order_style in style_xfty_groups:
            order_base = order_style.split('-')[0].split(' ')[0]
            if sheet_base == order_base:
                return order_style
        return None

    matched_styles = set()
    unmatched_styles = set(order_styles_set)

    FONT = 'Arial'
    normal_font = Font(name=FONT, size=10)
    thin = Side(style='thin', color='B0B0B0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center_alignment = openpyxl.styles.Alignment(horizontal='center', vertical='center')

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        order_key = find_order_key(sheet_name)
        if order_key is None:
            continue

        by_xfty = style_xfty_groups[order_key]
        matched_styles.add(order_key)
        if order_key in unmatched_styles:
            unmatched_styles.remove(order_key)

        sorted_dates = sorted(by_xfty.keys(), key=lambda d: (d is None, d or datetime.date.min))

        max_order_rows = max(len(sorted_dates), 10)
        for row_idx in range(4, 4 + max_order_rows):
            for col_idx in range(20, 29): # T to AB
                cell = ws.cell(row=row_idx, column=col_idx)
                if row_idx >= 4 + len(sorted_dates):
                    cell.value = None

        for i, x_fty_date in enumerate(sorted_dates):
            r = 4 + i
            ws.row_dimensions[r].height = 22
            data = by_xfty[x_fty_date]

            ws.cell(row=r, column=20, value=sheet_name) # T: Style
            
            u_cell = ws.cell(row=r, column=21) # U: X-fty
            if x_fty_date is not None:
                if isinstance(x_fty_date, (datetime.date, datetime.datetime)):
                    u_cell.value = datetime.datetime(x_fty_date.year, x_fty_date.month, x_fty_date.day)
                else:
                    u_cell.value = x_fty_date
                u_cell.number_format = 'dd-mm-yyyy'
            else:
                u_cell.value = ''

            ws.cell(row=r, column=22, value=data['qty'])     # V: Qty
            ws.cell(row=r, column=23, value=data['sewing'])  # W: Sewing Complete
            ws.cell(row=r, column=24, value=f'=W{r}-V{r}')  # X: Bal to sew (Completed - Qty)
            ws.cell(row=r, column=25, value=data['wash'])    # Y: Wash Complete
            ws.cell(row=r, column=26, value=f'=Y{r}-V{r}')  # Z: Bal to Wash (Completed - Qty)
            ws.cell(row=r, column=27, value=data['pack'])    # AA: Pack Complete
            ws.cell(row=r, column=28, value=f'=AA{r}-V{r}') # AB: Bal to Pack (Completed - Qty)

            for col_idx in range(20, 29):
                c = ws.cell(row=r, column=col_idx)
                c.font = normal_font
                c.alignment = center_alignment
                c.border = border

    try:
        wb.save(out_path)
    except Exception as e:
        fail(f"Could not save output file: {e}")

    print(json.dumps({
        "success": True,
        "stylesMatched": len(matched_styles),
        "stylesNotFound": len(unmatched_styles),
        "matchedStyles": sorted(matched_styles),
        "unmatchedStyles": sorted(unmatched_styles),
    }))


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        fail(f"Processing error: {e}")
