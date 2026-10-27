#!/usr/bin/env python3
"""
fill_wash_pack.py -- fill Wash & Pack Actuals into a previously-generated
Plan vs Actuals workbook using a Wash & Pack Daily Performance Report (.xlsx).

Usage:
    python3 fill_wash_pack.py <generated_workbook.xlsx> <wash_pack_report.xlsx> <output.xlsx>
"""
import sys
import os
import re
import json
import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment


def fail(message):
    print(json.dumps({"error": message}), file=sys.stderr)
    sys.exit(1)


MONTH_MAP = {
    'jan': 1, 'january': 1,
    'feb': 2, 'february': 2,
    'mar': 3, 'march': 3,
    'apr': 4, 'april': 4,
    'may': 5,
    'jun': 6, 'june': 6,
    'jul': 7, 'july': 7,
    'aug': 8, 'august': 8,
    'sep': 9, 'sept': 9, 'september': 9,
    'oct': 10, 'october': 10,
    'nov': 11, 'november': 11,
    'dec': 12, 'december': 12
}


def parse_date_from_string(text, ref_year=None):
    if not text:
        return None

    cleaned = os.path.splitext(os.path.basename(str(text)))[0]
    cleaned = re.sub(r'\(\d+\)', '', cleaned).strip()
    y_ref = ref_year or datetime.date.today().year

    # 1. Full date pattern: YYYY-MM-DD or YYYY_MM_DD
    m = re.search(r'(\d{4})[-_./](\d{1,2})[-_./](\d{1,2})', cleaned)
    if m:
        y, mth, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime.date(y, mth, d)
        except ValueError:
            pass

    # 2. Full date pattern: DD-MM-YYYY or DD_MM_YYYY
    m = re.search(r'(\d{1,2})[-_./](\d{1,2})[-_./](\d{4})', cleaned)
    if m:
        d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime.date(y, mth, d)
        except ValueError:
            pass

    # 3. Date with text month: DD-Mon-YYYY (e.g. 03-Sep-2026)
    m = re.search(r'(\d{1,2})[-_.\s]+([A-Za-z]{3,9})[-_.\s]+(\d{2,4})', cleaned)
    if m:
        d_str, m_str, y_str = m.group(1), m.group(2), m.group(3)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            y = int(y_str) if len(y_str) == 4 else 2000 + int(y_str)
            try:
                return datetime.date(y, mth, int(d_str))
            except ValueError:
                pass

    # 4. Mon-DD-YYYY or Mon DD YYYY
    m = re.search(r'([A-Za-z]{3,9})[-_.\s]+(\d{1,2})[-_.\s]+(\d{2,4})', cleaned)
    if m:
        m_str, d_str, y_str = m.group(1), m.group(2), m.group(3)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            y = int(y_str) if len(y_str) == 4 else 2000 + int(y_str)
            try:
                return datetime.date(y, mth, int(d_str))
            except ValueError:
                pass

    # 5. DD-Mon or DD Mon or DD_Mon (e.g. 3-sep-U1, 1-sep-UI, 5-sep-U1)
    m = re.search(r'(\d{1,2})[-_.\s]*([A-Za-z]{3,9})(?=[-_.\s]|$)', cleaned)
    if m:
        d_str, m_str = m.group(1), m.group(2)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            try:
                return datetime.date(y_ref, mth, int(d_str))
            except ValueError:
                pass

    # 6. Mon-DD or Mon DD or Mon_DD (e.g. SEP 15, SEP-15, DPR SEP 15)
    m = re.search(r'([A-Za-z]{3,9})[-_.\s]*(\d{1,2})(?=[-_.\s]|$)', cleaned)
    if m:
        m_str, d_str = m.group(1), m.group(2)
        mth = MONTH_MAP.get(m_str.lower())
        if mth:
            try:
                return datetime.date(y_ref, mth, int(d_str))
            except ValueError:
                pass

    # 7. Short date pattern: DD-MM or DD_MM (e.g. wash_pack_03_9)
    m = re.search(r'(\d{1,2})[-_./](\d{1,2})', cleaned)
    if m:
        d, mth = int(m.group(1)), int(m.group(2))
        try:
            return datetime.date(y_ref, mth, d)
        except ValueError:
            pass

    return None


def parse_number(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        s = val.replace(',', '').replace(' ', '').replace('$', '').strip()
        try:
            return float(s)
        except ValueError:
            return 0.0
    return 0.0


def extract_report_date(report_path, wb_perf, ref_year=None):
    d = parse_date_from_string(report_path, ref_year)
    if d:
        return d
    if wb_perf.sheetnames:
        d = parse_date_from_string(wb_perf.sheetnames[0], ref_year)
        if d:
            return d
    return None


def select_report_sheet(wb_perf, report_date, ref_year=None):
    """Pick the worksheet that actually matches report_date, instead of
    blindly using the first sheet.

    Real DPR exports from this factory's system are often *rolling-window*
    files that bundle 2-3 trailing dates in one workbook (e.g. a "05-Sep"
    export containing sheets '03-9-2026', '04-9-2026', '5-9-2026'). Sheet
    names also vary in zero-padding and sometimes carry a stray " (2)" suffix
    from being re-saved. Blindly taking worksheets[0] silently grabs a STALE,
    already-processed date instead of the new one.
    """
    for ws in wb_perf.worksheets:
        d = parse_date_from_string(ws.title, ref_year)
        if d == report_date:
            return ws, True

    # No sheet name matched the target date at all. These exports always
    # append the newest date last, so the last sheet is the best fallback --
    # far better than the first sheet, which is guaranteed to be old data.
    return wb_perf.worksheets[-1], False


def find_report_columns(ws_perf):
    rows = list(ws_perf.iter_rows(values_only=False))
    if not rows:
        return None, None, None, None, 0

    max_scan = min(10, len(rows))

    style_col = None
    wash_on_date_col = None
    pack_on_date_col = None
    po_col = None
    data_start_row = 1

    header_matrix = []
    for r_idx in range(max_scan):
        row_vals = [str(cell.value).strip() if cell and cell.value is not None else "" for cell in rows[r_idx]]
        header_matrix.append(row_vals)

    num_cols = max(len(r) for r in header_matrix) if header_matrix else 0

    for c in range(num_cols):
        for r in range(max_scan):
            val = header_matrix[r][c].upper()
            if val in ("STYLE", "STYLE NO", "STYLE NO.", "STYLE_NO", "STYLE NAME", "STYLENO") or (
                "STYLE" in val and "BUYER" not in val and "CUSTOMER" not in val and "LINE" not in val
            ):
                style_col = c
                break
        if style_col is not None:
            break

    for c in range(num_cols):
        for r in range(max_scan):
            val = header_matrix[r][c].upper().replace(" ", "").replace("_", "")
            if val in ("PO", "PONO", "BUYERPONO", "ORDERNO", "PONUMBER", "BUYERPO") or (
                "PONO" in val and "STYLE" not in val
            ):
                po_col = c
                break
        if po_col is not None:
            break

    # Identify Section Headers explicitly
    wash_sec_start = None
    pack_sec_start = None

    for r in range(max_scan):
        for c in range(num_cols):
            if c < len(header_matrix[r]):
                raw = header_matrix[r][c].upper()
                clean = raw.replace(" ", "").replace("_", "").replace("-", "")

                is_wash_recd = (
                    (("WASH" in clean or "W/REC" in clean) and any(k in clean for k in ("RECD", "REC", "RECEIVE", "RECEIVED")) and "SENT" not in clean)
                    or ("WASHING-RECD" in raw or "WASHING RECD" in raw or "WASH RECD" in raw or "W/REC" in raw)
                ) and "SENT" not in clean

                if is_wash_recd and wash_sec_start is None:
                    wash_sec_start = c

                is_packing = ("PACK" in clean or "PACKING" in clean or "PACKED" in clean) and "UNPACK" not in clean

                if is_packing and pack_sec_start is None:
                    pack_sec_start = c

    top_header_cols = set()
    for r in range(min(3, max_scan)):
        for c in range(num_cols):
            if c < len(header_matrix[r]) and header_matrix[r][c]:
                top_header_cols.add(c)
    sorted_top_cols = sorted(list(top_header_cols))

    def get_section_col_range(start_c):
        next_cols = [c for c in sorted_top_cols if c > start_c]
        end_c = (next_cols[0] - 1) if next_cols else (start_c + 3)
        return start_c, max(start_c, end_c)

    if wash_sec_start is not None:
        w_min, w_max = get_section_col_range(wash_sec_start)
        for c in range(w_min, min(w_max + 1, num_cols)):
            for r in range(max_scan):
                if c < len(header_matrix[r]):
                    cv = header_matrix[r][c].upper().replace(" ", "").replace("_", "").replace("-", "")
                    if cv in ("ONDATE", "TODAY", "ON", "DAILY"):
                        wash_on_date_col = c
                        break
            if wash_on_date_col is not None:
                break

    if pack_sec_start is not None:
        p_min, p_max = get_section_col_range(pack_sec_start)
        for c in range(p_min, min(p_max + 1, num_cols)):
            for r in range(max_scan):
                if c < len(header_matrix[r]):
                    cv = header_matrix[r][c].upper().replace(" ", "").replace("_", "").replace("-", "")
                    if cv in ("ONDATE", "TODAY", "ON", "DAILY"):
                        pack_on_date_col = c
                        break
            if pack_on_date_col is not None:
                break

    for r in range(max_scan):
        for c in range(num_cols):
            val = header_matrix[r][c].upper().replace(" ", "").replace("_", "")
            if "ONDATE" in val or val == "STYLE":
                data_start_row = max(data_start_row, r + 2)

    return style_col, wash_on_date_col, pack_on_date_col, po_col, data_start_row



IS_NOISE_EXACT = {
    "LINE-1", "LINE-2", "LINE-3", "LINE-4", "LINE-5", "LINE-6", "LINE-7", "LINE-8", "LINE-9", "LINE-10",
    "LINE-11", "LINE-12", "LINE-13", "LINE-13,5", "LINE-14", "LINE-15",
    "ENZYME+SOAPING", "NO WASH", "ENZYME WASH", "LIGHT BLEACH WASH", "HEAVY BLEACH WASH", "NORMAL WASH",
    "TOTAL", "GRAND TOTAL", "SUM", "STYLE", "STYLE NO", "STYLE NO.", "SL NO", "PO NO", "COLOR", "RECIPE"
}

NOISE_KEYWORDS = (
    "WASH", "ENZYME", "BLEACH", "TINT", "SOAPING", "DETERGENT", "RECIPE",
    "SPRAY", "WHISKER", "SANDING", "GRINDING", "DESTROY", "SCRAPING",
)

def is_noise_style(style_str):
    if not style_str:
        return True
    s = str(style_str).strip()
    u = s.upper()
    if u in IS_NOISE_EXACT:
        return True
    if u.startswith("LINE-") or u.startswith("LINE ") or u.startswith("RECIPE"):
        return True
    if u in ("DATE", "SL", "PO", "QUANTITY", "REMARKS"):
        return True
    if ("+" in s or " " in s) and any(k in u for k in NOISE_KEYWORDS):
        return True
    return False


def is_summary_row(row_cells, style_candidate=None):
    if style_candidate:
        sc = style_candidate.strip().upper()
        tokens = set(re.split(r'[\s_:\-]+', sc))
        if any(t in ("TOTAL", "SUBTOTAL", "GRANDTOTAL", "SUM", "TOTALS") for t in tokens):
            return True

    for cell in row_cells:
        if cell is None:
            continue
        c_str = str(cell).strip().upper()
        if c_str in ("TOTAL", "SUBTOTAL", "GRAND TOTAL", "SUM", "TOTALS"):
            return True
        tokens = set(re.split(r'[\s_:\-]+', c_str))
        if "SUBTOTAL" in tokens or "GRANDTOTAL" in tokens:
            return True
        if "TOTAL" in tokens and "STYLE" not in tokens and "NO" not in tokens:
            return True

    return False


def get_sheet_date(ws, report_file_date, ref_year=None):
    d = parse_date_from_string(ws.title, ref_year)
    if d:
        return d

    for row in list(ws.iter_rows(max_row=10, values_only=True)):
        for cell in row:
            if cell is not None:
                if isinstance(cell, datetime.datetime):
                    return cell.date()
                if isinstance(cell, datetime.date):
                    return cell
                sd = parse_date_from_string(str(cell), ref_year)
                if sd:
                    return sd

    return report_file_date


def load_excel_workbook(filepath, data_only=True):
    try:
        return openpyxl.load_workbook(filepath, data_only=data_only)
    except Exception as e:
        err_msg = str(e).lower()
        if "xls" in err_msg or filepath.lower().endswith(".xls"):
            try:
                import xlrd
            except ImportError:
                fail("Could not read .xls file: xlrd module is not installed.")

            try:
                wb_xls = xlrd.open_workbook(filepath)
            except Exception as ex_xls:
                fail(f"Could not open .xls file: {ex_xls}")

            wb_out = openpyxl.Workbook()
            wb_out.remove(wb_out.active)

            for sheet_name in wb_xls.sheet_names():
                sheet_xls = wb_xls.sheet_by_name(sheet_name)
                ws_out = wb_out.create_sheet(title=sheet_name)
                for r in range(sheet_xls.nrows):
                    row_vals = []
                    for c in range(sheet_xls.ncols):
                        cell_val = sheet_xls.cell_value(r, c)
                        cell_type = sheet_xls.cell_type(r, c)
                        if cell_type == xlrd.XL_CELL_DATE:
                            try:
                                dt_tuple = xlrd.xldate_as_tuple(cell_val, wb_xls.datemode)
                                if dt_tuple[0] == 0:
                                    cell_val = datetime.time(*dt_tuple[3:])
                                else:
                                    cell_val = datetime.datetime(*dt_tuple)
                            except Exception:
                                pass
                        elif cell_type in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            cell_val = None
                        elif cell_type == xlrd.XL_CELL_NUMBER:
                            if isinstance(cell_val, float) and cell_val.is_integer():
                                cell_val = int(cell_val)
                        row_vals.append(cell_val)
                    ws_out.append(row_vals)
            return wb_out
        else:
            raise e


def main():
    if len(sys.argv) < 3:
        fail("Usage: fill_wash_pack.py <generated_workbook.xlsx> <wash_pack_report.xlsx> <output.xlsx>")

    workbook_path = sys.argv[1]
    report_path = sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) >= 4 else "output.xlsx"

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except Exception as e:
        fail(f"Could not open the generated workbook: {e}")

    ref_year = None
    for ws in wb.worksheets:
        for r in range(4, min(15, ws.max_row + 1)):
            v = ws.cell(row=r, column=3).value
            if v and hasattr(v, "year"):
                ref_year = v.year
                break
        if ref_year:
            break

    try:
        wb_perf = load_excel_workbook(report_path, data_only=True)
    except Exception as e:
        fail(f"Could not open the wash & pack report: {e}")

    report_file_date = parse_date_from_string(report_path, ref_year)

    sheet_names = wb.sheetnames

    combined_bases = set()
    for name in sheet_names:
        if name.endswith("-Jacket"):
            base = name[:-len("-Jacket")]
            if f"{base}-Pant" in sheet_names:
                combined_bases.add(base)

    def resolve_targets(raw_style):
        if raw_style in sheet_names:
            return [("direct", raw_style)]

        for sn in sheet_names:
            if "&" in sn:
                parts = [p.strip() for p in sn.split("&")]
                if raw_style in parts:
                    return [("direct", sn)]

        if raw_style in combined_bases or raw_style.replace("-SET", "").replace("_SET", "") in combined_bases:
            base = raw_style.replace("-SET", "").replace("_SET", "")
            return [("combined", base)]

        for sn in sheet_names:
            if raw_style == sn or raw_style.startswith(sn + "-") or sn.startswith(raw_style + "-") or raw_style.startswith(sn + " "):
                return [("direct", sn)]

        return []

    wash_actuals = defaultdict(float)  # key: (date_obj, style_key) -> sum
    pack_actuals = defaultdict(float)  # key: (date_obj, style_key) -> sum
    date_matched_has_rows = defaultdict(lambda: defaultdict(bool))  # date_obj -> style_key -> bool
    unmatched_styles_by_date = defaultdict(set)
    matched_styles_by_date = defaultdict(set)
    processed_sheets_info = []

    # Every date that was actually present/readable in the uploaded Wash &
    # Pack report (i.e. we found a usable sheet + style/wash/pack columns
    # for it). This is what tells us the upload "covers" that date -- so a
    # style with no rows for it should get a real 0, not a blank. Dates
    # that never appear here simply were not uploaded, and must stay blank.
    report_dates = set()

    for ws_perf in wb_perf.worksheets:
        sheet_date = get_sheet_date(ws_perf, report_file_date, ref_year)
        if not sheet_date:
            continue

        style_col, wash_col, pack_col, po_col, data_start_row = find_report_columns(ws_perf)

        if style_col is None or (wash_col is None and pack_col is None):
            continue

        report_dates.add(sheet_date)

        styles_detected = set()
        wash_values_count = 0
        pack_values_count = 0
        skipped_totals_count = 0
        skipped_invalid_count = 0

        perf_rows = list(ws_perf.iter_rows(values_only=True))
        current_style = None

        for r_idx in range(data_start_row - 1, len(perf_rows)):
            r = perf_rows[r_idx]
            if not r:
                skipped_invalid_count += 1
                continue

            raw_cell = r[style_col] if len(r) > style_col else None
            style_candidate = str(raw_cell).strip() if (raw_cell is not None and str(raw_cell).strip()) else None

            if is_summary_row(r, style_candidate):
                skipped_totals_count += 1
                continue

            if style_candidate is not None:
                if not is_noise_style(style_candidate):
                    current_style = style_candidate

            if current_style is None:
                skipped_invalid_count += 1
                continue

            if po_col is not None:
                po_val = r[po_col] if len(r) > po_col else None
                if po_val is None or str(po_val).strip() == "":
                    skipped_totals_count += 1
                    continue

            w_val = parse_number(r[wash_col]) if (wash_col is not None and len(r) > wash_col) else 0.0
            p_val = parse_number(r[pack_col]) if (pack_col is not None and len(r) > pack_col) else 0.0

            if w_val == 0.0 and p_val == 0.0:
                continue

            styles_detected.add(current_style)
            if w_val > 0:
                wash_values_count += 1
            if p_val > 0:
                pack_values_count += 1

            targets = resolve_targets(current_style)
            if not targets:
                unmatched_styles_by_date[str(sheet_date)].add(current_style)
                continue

            for kind, key in targets:
                matched_styles_by_date[str(sheet_date)].add(key)
                if kind == "direct":
                    wash_actuals[(sheet_date, key)] += w_val
                    pack_actuals[(sheet_date, key)] += p_val
                    date_matched_has_rows[sheet_date][key] = True
                else:
                    jacket_sheet = f"{key}-Jacket"
                    pant_sheet = f"{key}-Pant"
                    wash_actuals[(sheet_date, jacket_sheet)] += round(w_val / 2)
                    wash_actuals[(sheet_date, pant_sheet)] += w_val - round(w_val / 2)
                    pack_actuals[(sheet_date, jacket_sheet)] += round(p_val / 2)
                    pack_actuals[(sheet_date, pant_sheet)] += p_val - round(p_val / 2)
                    date_matched_has_rows[sheet_date][jacket_sheet] = True
                    date_matched_has_rows[sheet_date][pant_sheet] = True

        def col_letter(idx):
            return openpyxl.utils.get_column_letter(idx + 1) if idx is not None else "N/A"

        info = {
            "sheetTitle": ws_perf.title,
            "detectedDate": str(sheet_date),
            "styleCol": col_letter(style_col),
            "washCol": col_letter(wash_col),
            "packCol": col_letter(pack_col),
            "stylesDetected": len(styles_detected),
            "washValuesFound": wash_values_count,
            "packValuesFound": pack_values_count,
            "skippedTotals": skipped_totals_count
        }
        processed_sheets_info.append(info)

        print(f"Sheet: {ws_perf.title}", file=sys.stderr)
        print(f"  Detected date: {sheet_date}", file=sys.stderr)
        print(f"  Style column: {col_letter(style_col)}", file=sys.stderr)
        print(f"  Wash ON DATE column: {col_letter(wash_col)}", file=sys.stderr)
        print(f"  Pack ON DATE column: {col_letter(pack_col)}", file=sys.stderr)
        print(f"  Styles detected: {len(styles_detected)}", file=sys.stderr)
        print(f"  Wash values found: {wash_values_count}", file=sys.stderr)
        print(f"  Pack values found: {pack_values_count}", file=sys.stderr)
        print(f"  Skipped TOTAL rows: {skipped_totals_count}", file=sys.stderr)

    if not processed_sheets_info:
        fail("No valid DPR sheets found in Wash & Pack report.")

    sheets_updated = 0
    cells_filled = 0
    center_align = Alignment(horizontal='center', vertical='center')

    valid_report_dates = [rd for rd in report_dates if isinstance(rd, datetime.date)]
    max_report_date = max(valid_report_dates) if valid_report_dates else None

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        touched = False
        last_date_row = 4

        for r in range(4, ws.max_row + 1):
            date_val = ws.cell(row=r, column=3).value
            if date_val is None:
                continue

            last_date_row = r
            parsed_d = parse_date_from_string(date_val, ref_year)
            d = parsed_d if parsed_d is not None else (date_val.date() if hasattr(date_val, "date") else date_val)

            key = (d, sheet_name)
            has_data = (key in wash_actuals) or (key in pack_actuals) or date_matched_has_rows[d][sheet_name]

            if has_data:
                w_tot = wash_actuals[key]
                p_tot = pack_actuals[key]

                k_cell = ws.cell(row=r, column=11, value=int(w_tot) if w_tot.is_integer() else w_tot)
                p_cell = ws.cell(row=r, column=16, value=int(p_tot) if p_tot.is_integer() else p_tot)
                k_cell.alignment = center_align
                p_cell.alignment = center_align
                cells_filled += 2
                touched = True
            elif d in report_dates:
                # This date was covered by the uploaded Wash & Pack report,
                # but this particular style had no wash/pack rows in it --
                # that's a genuine 0 for the day, not a blank.
                k_cell = ws.cell(row=r, column=11, value=0)
                p_cell = ws.cell(row=r, column=16, value=0)
                k_cell.alignment = center_align
                p_cell.alignment = center_align
                cells_filled += 2
                touched = True
            elif max_report_date and isinstance(d, datetime.date) and d > max_report_date:
                # Date is past the maximum uploaded report date -> clear actuals
                ws.cell(row=r, column=11, value=None)
                ws.cell(row=r, column=16, value=None)

            # Preserve & re-assert exact IF(ISNUMBER(...)) formulas across ALL 3 blocks for every row 4..max_row
            ws.cell(row=r, column=5, value=f'=IF(ISNUMBER(F{r}),SUM($D$4:D{r}),"")')     # E: Sew Cum Plan
            ws.cell(row=r, column=7, value=f'=IF(ISNUMBER(F{r}),SUM($F$4:F{r}),"")')     # G: Sew Cum Actual
            ws.cell(row=r, column=8, value=f'=IF(ISNUMBER(G{r}),N(E{r})-G{r},"")')      # H: Sew Deviation (blank Cum Plan treated as 0)
            ws.cell(row=r, column=10, value=f'=IF(ISNUMBER(K{r}),SUM($I$4:I{r}),"")')   # J: Wash Cum Plan
            ws.cell(row=r, column=12, value=f'=IF(ISNUMBER(K{r}),SUM($K$4:K{r}),"")')   # L: Wash Cum Actual
            ws.cell(row=r, column=13, value=f'=IF(ISNUMBER(L{r}),N(J{r})-L{r},"")')     # M: Wash Deviation (blank Cum Plan treated as 0)
            ws.cell(row=r, column=15, value=f'=IF(ISNUMBER(P{r}),SUM($N$4:N{r}),"")')   # O: Pack Cum Plan
            ws.cell(row=r, column=17, value=f'=IF(ISNUMBER(P{r}),SUM($P$4:P{r}),"")')   # Q: Pack Cum Actual
            ws.cell(row=r, column=18, value=f'=IF(ISNUMBER(Q{r}),N(O{r})-Q{r},"")')     # R: Pack Deviation (blank Cum Plan treated as 0)

            for col_idx in (5, 7, 8, 10, 12, 13, 15, 17, 18):
                ws.cell(row=r, column=col_idx).alignment = center_align

        # NOTE: Order Summary columns (T to AB) are intentionally NOT touched here.
        # That section (Style, X-fty, Qty, Sewing completed, Bal to sew, Wash
        # completed, Bal to Wash, Pack completed, Bal to Pack) is maintained
        # manually by the user in the uploaded Existing Workbook and must be
        # returned exactly as uploaded -- this script only fills the Sewing
        # Details / Wash Details / Pack Details block (columns A-R).

        if touched:
            sheets_updated += 1

    try:
        wb.save(out_path)
    except Exception as e:
        fail(f"Could not save output file: {e}")

    all_unmatched = sorted(set().union(*unmatched_styles_by_date.values()))

    print(json.dumps({
        "success": True,
        "processedSheets": processed_sheets_info,
        "sheetsUpdated": sheets_updated,
        "cellsFilled": cells_filled,
        "unmatchedStyles": all_unmatched,
    }))


if __name__ == "__main__":
    main()