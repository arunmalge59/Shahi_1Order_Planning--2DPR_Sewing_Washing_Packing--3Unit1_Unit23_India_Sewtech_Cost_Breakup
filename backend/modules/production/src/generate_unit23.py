#!/usr/bin/env python3
"""
generate_unit23.py -- Book1-style line plan -> Style-wise Plan vs Actuals workbook for Unit 23

Flexible header parsing to support Unit 23 Book1 line plans with column header variations:
- Style column: "style", "style number", "style_number", "style no", "style no.", "style name", "styleno", "styles"
- Item column: "item no.", "item no", "item", "item_no", "itemcode", "item code", "article", "cat", "category", "desc", "description" (optional fallback if omitted)
- Contract Qty column: "contract qty", "contract quantity", "contract", "order qty", "order quantity", "qty", "total qty" (optional fallback)
- Date columns: Datetime objects, Date objects, Excel serial dates, or formatted strings ("17'Apr", "30-Mar", "2026-09-01", "1-Sep", etc.)
"""
import sys
import os
import re
import json
import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.formatting.rule import FormulaRule


def fail(message):
    print(json.dumps({"error": message}), file=sys.stderr)
    sys.exit(1)


def parse_header_date(h, ref_year=None):
    if h is None:
        return None
    if isinstance(h, datetime.datetime):
        return h.date()
    if isinstance(h, datetime.date):
        return h
    if isinstance(h, (int, float)):
        if 40000 <= h <= 60000:
            try:
                return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(h))).date()
            except Exception:
                return None
        return None

    s = str(h).strip()
    if not s:
        return None

    y_ref = ref_year or datetime.date.today().year

    # Match DD'Mon or DD-Mon or DD_Mon or DD Mon (e.g. 17'Apr, 30'Mar, 1-Sep)
    m = re.search(r"^(\d{1,2})['_\-\s/]([A-Za-z]{3,9})$", s)
    if m:
        d_str, m_str = m.group(1), m.group(2)
        for fmt in ("%d-%b-%Y", "%d-%B-%Y"):
            try:
                return datetime.datetime.strptime(f"{d_str}-{m_str}-{y_ref}", fmt).date()
            except ValueError:
                continue

    # Match Mon'DD or Mon-DD
    m = re.search(r"^([A-Za-z]{3,9})['_\-\s/](\d{1,2})$", s)
    if m:
        m_str, d_str = m.group(1), m.group(2)
        for fmt in ("%b-%d-%Y", "%B-%d-%Y"):
            try:
                return datetime.datetime.strptime(f"{m_str}-{d_str}-{y_ref}", fmt).date()
            except ValueError:
                continue

    # Standard formats
    for fmt in (
        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y",
        "%d-%b-%Y", "%d-%b-%y", "%d-%m-%y", "%Y/%m/%d"
    ):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue

    return None


def main():
    if len(sys.argv) != 3:
        fail("Usage: generate_unit23.py <input.xlsx> <output.xlsx>")

    src, out = sys.argv[1], sys.argv[2]

    try:
        wb_in = openpyxl.load_workbook(src, data_only=True, read_only=True)
    except Exception as e:
        fail(f"Could not open input file: {e}")


    # Search across sheets for the first sheet with data
    ws_in = None
    for ws in wb_in.worksheets:
        if ws.max_row > 1:
            ws_in = ws
            break

    if ws_in is None:
        ws_in = wb_in.worksheets[0]

    rows = list(ws_in.iter_rows(values_only=True))
    if not rows:
        fail("Input sheet is empty.")

    header_row_idx = None
    for idx, r in enumerate(rows[:30]):
        if not r:
            continue
        cleaned = [str(c).strip().lower() if c is not None else '' for c in r]
        if any('style' in c for c in cleaned):
            header_row_idx = idx
            break

    if header_row_idx is None:
        fail("Input file is missing a 'Style' column header.")

    header = rows[header_row_idx]
    cleaned_header_lower = [str(c).strip().lower() if c is not None else '' for c in header]

    # Flexible Style column search
    style_idx = None
    for i, c in enumerate(cleaned_header_lower):
        if 'style' in c:
            style_idx = i
            break

    if style_idx is None:
        fail("Input file is missing a 'Style' column.")

    # Flexible Item column search
    item_idx = None
    for candidate in ('item no.', 'item no', 'item', 'item_no', 'itemcode', 'item code', 'article', 'cat', 'category', 'desc', 'description'):
        if candidate in cleaned_header_lower:
            item_idx = cleaned_header_lower.index(candidate)
            break
    if item_idx is None:
        for i, c in enumerate(cleaned_header_lower):
            if any(k in c for k in ('item', 'article', 'cat', 'desc')):
                item_idx = i
                break

    # Flexible Contract Qty column search
    contract_idx = None
    for candidate in ('contract qty', 'contract quantity', 'contract', 'order qty', 'order quantity', 'qty', 'total qty'):
        if candidate in cleaned_header_lower:
            contract_idx = cleaned_header_lower.index(candidate)
            break

    # Parse date columns
    date_col_idx = []
    dates = []
    ref_year = datetime.date.today().year

    for i, h in enumerate(header):
        parsed_d = parse_header_date(h, ref_year=ref_year)
        if parsed_d is not None:
            date_col_idx.append(i)
            dates.append(parsed_d)

    # Deduplicate and sort dates chronologically
    unique_dates = sorted(list(set(dates)))

    if not unique_dates:
        fail("No valid daily date columns found in the header row.")

    data_rows = []
    for r in rows[header_row_idx + 1:]:
        if not r or len(r) <= style_idx or r[style_idx] is None:
            continue
        st_val = str(r[style_idx]).strip()
        if not st_val or any(k in st_val.upper() for k in ("TOTAL", "SUBTOTAL", "GRAND TOTAL")):
            continue
        data_rows.append(r)

    if not data_rows:
        fail("No data rows found under the header.")

    style_item = {}
    style_daily = defaultdict(lambda: defaultdict(int))
    style_contract_qty = defaultdict(int)

    for r in data_rows:
        style = str(r[style_idx]).strip()
        item = str(r[item_idx]).strip() if (item_idx is not None and len(r) > item_idx and r[item_idx] is not None) else ""
        if style not in style_item:
            style_item[style] = item

        if contract_idx is not None and len(r) > contract_idx and r[contract_idx] is not None:
            try:
                style_contract_qty[style] += int(r[contract_idx])
            except (ValueError, TypeError):
                pass

        for pos, ci in enumerate(date_col_idx):
            v = r[ci] if len(r) > ci and r[ci] is not None else 0
            try:
                style_daily[style][dates[pos]] += int(v)
            except (ValueError, TypeError):
                pass

    styles = list(style_item.keys())
    dates = unique_dates


    # Build output workbook
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    FONT = 'Arial'
    header_font = Font(name=FONT, bold=True, size=10)
    group_font = Font(name=FONT, bold=True, size=11, color='FFFFFF')
    normal_font = Font(name=FONT, size=10)
    thin = Side(style='thin', color='B0B0B0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    group_fills = {
        'sew': PatternFill('solid', fgColor='0D9488'),
        'wash': PatternFill('solid', fgColor='2563EB'),
        'pack': PatternFill('solid', fgColor='C2410C'),
    }
    manual_fill = PatternFill('solid', fgColor='FFF2CC')
    deviation_positive_fill = PatternFill('solid', fgColor='C6EFCE')
    deviation_positive_font = Font(name=FONT, size=10, color='006100')
    deviation_negative_fill = PatternFill('solid', fgColor='FFC7CE')
    deviation_negative_font = Font(name=FONT, size=10, color='9C0006')

    def build_sheet(wb, style):
        ws = wb.create_sheet(title=str(style)[:31])
        item = style_item[style]

        ws.merge_cells('D2:H2'); ws['D2'] = 'Sewing Details'
        ws.merge_cells('I2:M2'); ws['I2'] = 'Wash Details'
        ws.merge_cells('N2:R2'); ws['N2'] = 'Pack Details'
        center_alignment = Alignment(horizontal='center', vertical='center')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        ws.row_dimensions[2].height = 26
        ws.row_dimensions[3].height = 26

        for cellref, key in (('D2', 'sew'), ('I2', 'wash'), ('N2', 'pack')):
            c = ws[cellref]
            c.font = group_font
            c.alignment = center_alignment
            c.fill = group_fills[key]

        cols = ['Style', 'Item No.', 'Date',
                'Plan Qty', 'Cum Plan', 'Actuals', 'Cum Actual', 'Deviation',
                'Plan Qty', 'Cum Plan', 'Actuals', 'Cum Actual', 'Deviation',
                'Plan Qty', 'Cum Plan', 'Actuals', 'Cum Actual', 'Deviation']
        for j, name in enumerate(cols, start=1):
            c = ws.cell(row=3, column=j, value=name)
            c.font = header_font
            c.alignment = header_alignment
            c.border = border

        sew_plan_prev = None
        wash_plan_prev = None

        for i, d in enumerate(dates):
            r = 4 + i
            ws.row_dimensions[r].height = 22
            sew_plan = style_daily[style][d]
            wash_plan = sew_plan_prev
            pack_plan = wash_plan_prev

            ws.cell(row=r, column=1, value=style)
            ws.cell(row=r, column=2, value=item)
            dc = ws.cell(row=r, column=3, value=datetime.datetime(d.year, d.month, d.day))
            dc.number_format = 'dd-mm-yyyy'

            ws.cell(row=r, column=4, value=sew_plan)
            ws.cell(row=r, column=5, value=f'=IF(ISNUMBER(F{r}),SUM($D$4:D{r}),"")')
            ws.cell(row=r, column=6).fill = manual_fill
            ws.cell(row=r, column=7, value=f'=IF(ISNUMBER(F{r}),SUM($F$4:F{r}),"")')
            ws.cell(row=r, column=8, value=f'=IF(ISNUMBER(G{r}),N(E{r})-G{r},"")')

            ws.cell(row=r, column=9, value=wash_plan)
            ws.cell(row=r, column=10, value=f'=IF(ISNUMBER(K{r}),SUM($I$4:I{r}),"")')
            ws.cell(row=r, column=11).fill = manual_fill
            ws.cell(row=r, column=12, value=f'=IF(ISNUMBER(K{r}),SUM($K$4:K{r}),"")')
            ws.cell(row=r, column=13, value=f'=IF(ISNUMBER(L{r}),N(J{r})-L{r},"")')

            ws.cell(row=r, column=14, value=pack_plan)
            ws.cell(row=r, column=15, value=f'=IF(ISNUMBER(P{r}),SUM($N$4:N{r}),"")')
            ws.cell(row=r, column=16).fill = manual_fill
            ws.cell(row=r, column=17, value=f'=IF(ISNUMBER(P{r}),SUM($P$4:P{r}),"")')
            ws.cell(row=r, column=18, value=f'=IF(ISNUMBER(Q{r}),N(O{r})-Q{r},"")')

            for j in range(1, 19):
                cell = ws.cell(row=r, column=j)
                cell.font = normal_font
                cell.alignment = center_alignment
                cell.border = border

            sew_plan_prev = sew_plan
            wash_plan_prev = wash_plan

        for column in ('H', 'M', 'R'):
            deviation_range = f'{column}4:{column}{3 + len(dates)}'
            ws.conditional_formatting.add(
                deviation_range,
                FormulaRule(
                    formula=[f'{column}4>0'],
                    fill=deviation_positive_fill,
                    font=deviation_positive_font,
                ),
            )
            ws.conditional_formatting.add(
                deviation_range,
                FormulaRule(
                    formula=[f'{column}4<0'],
                    fill=deviation_negative_fill,
                    font=deviation_negative_font,
                ),
            )

        ws['T3'] = 'Style'; ws['U3'] = 'X-fty'; ws['V3'] = 'Qty'
        ws['W3'] = 'Sewing completd'; ws['X3'] = 'Bal to sew'
        ws['Y3'] = 'Wash completd'; ws['Z3'] = 'Bal to Wash'
        ws['AA3'] = 'Pack completd'; ws['AB3'] = 'Bal to Pack'
        for col in ['T', 'U', 'V', 'W', 'X', 'Y', 'Z', 'AA', 'AB']:
            c = ws[f'{col}3']
            c.font = header_font
            c.alignment = header_alignment
            c.border = border
        ws['T4'] = style
        ws['V4'] = style_contract_qty.get(style, '')
        for col in ['T', 'U', 'V', 'W', 'X', 'Y', 'Z', 'AA', 'AB']:
            c = ws[f'{col}4']
            c.font = normal_font
            c.alignment = center_alignment
            c.border = border

        col_widths = {
            'A': 22, 'B': 14, 'C': 14,
            'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 12,
            'I': 12, 'J': 12, 'K': 12, 'L': 12, 'M': 12,
            'N': 12, 'O': 12, 'P': 12, 'Q': 12, 'R': 12,
            'S': 4,
            'T': 22, 'U': 14, 'V': 12, 'W': 16, 'X': 14,
            'Y': 16, 'Z': 14, 'AA': 16, 'AB': 14
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

    for style in styles:
        build_sheet(wb_out, style)

    try:
        wb_out.save(out)
    except Exception as e:
        fail(f"Could not save output file: {e}")

    print(json.dumps({
        "success": True,
        "styles": len(styles),
        "dates": len(dates),
        "rows": len(data_rows),
        "availableDates": [d.isoformat() for d in dates],
    }))


if __name__ == '__main__':
    main()
